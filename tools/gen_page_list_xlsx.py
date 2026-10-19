"""Build the vendor page list (Excel) from the website's own registry and English content."""
import json, sys, glob
from pathlib import Path
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / 'src'))
import registry as R
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

BASE = 'https://hytengineering.com'
# ---------- content (English master) ----------
C, common = {}, {}
for f in glob.glob(str(ROOT / 'src/content/en/*.json')):
    d = json.load(open(f))
    if Path(f).stem == 'common':
        common = d
    for k, v in d.get('pages', {}).items():
        C[k] = v
if 'home' not in C:
    C['home'] = json.load(open(ROOT / 'src/content/en/home.json'))
menu_t = common['menu']

def nav(pid):
    p = C.get(pid, {})
    return p.get('nav') or p.get('h1') or pid

PG = {p['id']: p for p in R.P}
kids = {}
for p in R.P:
    kids.setdefault(p['parent'], []).append(p)

# order children: follow menu order where possible, then sr
menu_order = []
for m in R.MENU:
    if m.get('id'): menu_order.append(m['id'])
    for c in m['columns']:
        for it in c['items']:
            menu_order.append(it[0] if isinstance(it, tuple) else it)
in_menu = set(menu_order)
rank = {pid: i for i, pid in enumerate(menu_order)}
for k in kids:
    kids[k].sort(key=lambda p: (rank.get(p['id'], 1000), p['sr']))

footer_ids = {i for col in R.FOOTER for i in col['items']} | set(R.LEGAL_FOOTER)

# which main-menu section a page lives under
section_of = {}
for m in R.MENU:
    label = nav(m['id']) if m.get('id') else menu_t.get(m.get('label') or 'media', 'Media & Resources')
    ids = ([m['id']] if m.get('id') else []) + [it[0] if isinstance(it, tuple) else it for c in m['columns'] for it in c['items']]
    for i in ids:
        section_of.setdefault(i, label)
section_of['contact-us'] = 'Contact Us (header button)'

TYPE = {'home': 'Home page', 'product': 'Product page', 'legal': 'Legal page', 'utility': 'Utility page', 'page': 'Content page'}
FEATURE = {'form': 'Form', 'facilities': 'Interactive tabs', 'faq': 'FAQ accordion', 'gallery': 'Image gallery + lightbox',
           'specs': 'Specification table', 'logo_carousel': 'Customer logo carousel', 'product_catalog': 'Product filter',
           'sitemap': 'Searchable sitemap', 'downloads': 'Download list', 'projects': 'Project cards', 'timeline': 'Timeline',
           'contact': 'Contact details + map link', 'news': 'News cards', 'cards': None, 'stats': 'Animated counters', 'steps': None}

def features(pid):
    p, t = C.get(pid, {}), PG[pid]['template']
    out = []
    for b in p.get('blocks', []):
        f = FEATURE.get(b['type'])
        if b['type'] == 'form':
            f = {'enquiry': 'Enquiry form', 'career': 'Career application form', 'supplier': 'Supplier registration form',
                 'catalogue': 'Catalogue request form'}.get(b.get('form', 'enquiry'), 'Form')
        if f and f not in out: out.append(f)
    if t == 'product':
        out = ['Product video', 'Catalogue download (form-gated)', 'On-page sub-navigation'] + out
    if p.get('hero_actions'): out.insert(0, 'Tall hero with action buttons')
    if any(b['type'] == 'notice' for b in p.get('blocks', [])): pass
    return ', '.join(out)

def sections(pid):
    p = C.get(pid, {})
    hs = [b.get('heading') for b in p.get('blocks', []) if b.get('heading') and b['type'] not in ('notice', 'cta')]
    clean = lambda s: s.replace('**', '')
    return ' | '.join(clean(h) for h in hs)

def pending(pid):
    p = C.get(pid, {})
    return 'Yes' if any(b['type'] == 'notice' for b in p.get('blocks', [])) else 'No'

rows = []
LEVEL = {0: 'Level 0 – Home', 1: 'Level 1 – Main page', 2: 'Level 2 – Sub page', 3: 'Level 3 – Sub-sub page', 4: 'Level 4'}
def walk(pid, depth):
    p = PG[pid]
    parent = p['parent']
    main = section_of.get(pid)
    if not main:
        anc = pid
        while PG[anc]['parent'] not in (None, 'home'): anc = PG[anc]['parent']
        if depth == 0: main = 'Home'
        elif anc in section_of: main = section_of[anc]
        elif p['template'] == 'legal': main = 'Footer – Legal'
        else: main = 'Footer / Utility'
    rows.append(dict(level=LEVEL[min(depth, 4)], depth=depth, main=main, parent=nav(parent) if parent else '—',
                     name=('    ' * max(depth - 1, 0)) + nav(pid), plain=nav(pid), url=BASE + p['path'],
                     ptype=TYPE.get(p['template'], 'Content page'), h1=C.get(pid, {}).get('h1', '').replace('**', ''),
                     sections=sections(pid), features=features(pid),
                     menu='Yes' if pid in in_menu else 'No', footer='Yes' if pid in footer_ids else 'No',
                     sitemap='Yes' if p.get('sitemap', True) else 'No (noindex)', pending=pending(pid)))
    for k in kids.get(pid, []):
        walk(k['id'], depth + 1)
walk('home', 0)
assert len(rows) == len(R.P), (len(rows), len(R.P))

# ---------- workbook ----------
wb = Workbook()
FONT = 'Arial'
H = Font(name=FONT, bold=True, color='FFFFFF', size=10)
HF = PatternFill('solid', fgColor='184A72')
B = Font(name=FONT, size=10)
BB = Font(name=FONT, size=10, bold=True)
thin = Side(style='thin', color='D0D5DB')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
LVL_FILL = {0: 'F58220', 1: 'DCE8F3', 2: 'FFFFFF', 3: 'F7F9FB', 4: 'F7F9FB'}
WRAP = Alignment(wrap_text=True, vertical='top')

def header(ws, cols, widths, row=1):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font, cell.fill, cell.border = H, HF, BOX
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row=row + 1, column=1)

# Sheet 1: Page List
ws = wb.active
ws.title = 'Page List'
cols = ['Sr. No.', 'Level', 'Main Menu Section', 'Parent Page', 'Page Name (hierarchy)', 'URL (English – Global)',
        'Page Type', 'Main Heading (H1)', 'Key Sections on Page', 'Special Features / Functionality',
        'In Main Menu', 'In Footer', 'In XML Sitemap', 'Content Pending HYT Approval', 'Vendor Remarks']
header(ws, cols, [8, 22, 22, 26, 44, 52, 16, 50, 70, 46, 11, 10, 14, 15, 30])
for i, r in enumerate(rows, 1):
    vals = [i, r['level'], r['main'], r['parent'], r['name'], r['url'], r['ptype'], r['h1'], r['sections'], r['features'],
            r['menu'], r['footer'], r['sitemap'], r['pending'], '']
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=i + 1, column=j, value=v)
        c.font = BB if r['depth'] <= 1 and j in (2, 5) else B
        if r['depth'] == 0 and j in (2, 5): c.font = Font(name=FONT, size=10, bold=True, color='FFFFFF')
        c.alignment = WRAP
        c.border = BOX
        c.fill = PatternFill('solid', fgColor=LVL_FILL[min(r['depth'], 4)])
    ws.cell(row=i + 1, column=6).hyperlink = r['url']
ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"
last = len(rows) + 1

# Sheet 2: Menu Structure
m2 = wb.create_sheet('Menu Structure')
header(m2, ['Top Menu Item', 'Dropdown Group', 'Menu Link', 'Indent', 'URL'], [24, 28, 46, 9, 56])
r = 2
for m in R.MENU:
    top = nav(m['id']) if m.get('id') else menu_t.get(m.get('label') or 'media', 'Media & Resources')
    for c in m['columns']:
        grp = menu_t.get(c.get('heading'), '') if c.get('heading') else '—'
        for it in c['items']:
            pid, ind = (it[0], 'Sub-item') if isinstance(it, tuple) else (it, '')
            for j, v in enumerate([top, grp, nav(pid), ind, BASE + PG[pid]['path']], 1):
                cell = m2.cell(row=r, column=j, value=v); cell.font = B; cell.border = BOX; cell.alignment = WRAP
            r += 1
for j, v in enumerate(['Contact Us (button)', '—', nav('contact-us'), '', BASE + PG['contact-us']['path']], 1):
    cell = m2.cell(row=r, column=j, value=v); cell.font = B; cell.border = BOX
r += 2
m2.cell(row=r, column=1, value='Top bar quick links').font = BB
for q in [('Knowledge', 'knowledge-centre'), ('Join our team', 'career'), ('Next Event (highlighted)', 'international-exhibitions'), ('Search (site search overlay)', 'sitemap')]:
    r += 1
    m2.cell(row=r, column=1, value=q[0]).font = B
    m2.cell(row=r, column=3, value=nav(q[1])).font = B
    m2.cell(row=r, column=5, value=BASE + PG[q[1]]['path']).font = B

# Sheet 3: Footer
f3 = wb.create_sheet('Footer Links')
header(f3, ['Footer Column', 'Link', 'URL'], [26, 46, 56])
r = 2
for col in R.FOOTER:
    for pid in col['items']:
        for j, v in enumerate([common['footer'].get(col['heading'], col['heading']), nav(pid), BASE + PG[pid]['path']], 1):
            cell = f3.cell(row=r, column=j, value=v); cell.font = B; cell.border = BOX
        r += 1
for pid in R.LEGAL_FOOTER:
    for j, v in enumerate(['Bottom strip – Legal', nav(pid), BASE + PG[pid]['path']], 1):
        cell = f3.cell(row=r, column=j, value=v); cell.font = B; cell.border = BOX
    r += 1
for j, v in enumerate(['Social media', 'LinkedIn', 'https://www.linkedin.com/company/hyt-engineering-company-private-limited/'], 1):
    f3.cell(row=r, column=j, value=v).font = B
r += 1
for j, v in enumerate(['Social media', 'YouTube', 'https://www.youtube.com/@HYTEngineeringCompanyPvtLtd/videos'], 1):
    f3.cell(row=r, column=j, value=v).font = B

# Sheet 4: Languages
l4 = wb.create_sheet('Languages')
header(l4, ['Language / Region', 'URL Folder', 'hreflang', 'Example URL', 'Pages (formula)', 'Translation'], [26, 14, 11, 52, 15, 40])
langs = [('English (Global) – default', '/', 'en', 'Master copy, British spelling'),
         ('German (Germany)', '/de-de/', 'de-DE', 'Translated – native review pending'),
         ('French (France)', '/fr-fr/', 'fr-FR', 'Translated – native review pending'),
         ('Italian (Italy)', '/it-it/', 'it-IT', 'Translated – native review pending'),
         ('Dutch (Netherlands)', '/nl-nl/', 'nl-NL', 'Translated – native review pending'),
         ('English (North America)', '/en-us/', 'en-US', 'Derived from English, American spelling')]
for i, (n, f, h, t) in enumerate(langs, 2):
    ex = BASE + (f if f != '/' else '/') + ('products/' if True else '')
    for j, v in enumerate([n, f, h, ex, f"=COUNTA('Page List'!A2:A{last})", t], 1):
        c = l4.cell(row=i, column=j, value=v); c.font = B; c.border = BOX
tr = len(langs) + 2
l4.cell(row=tr, column=1, value='Total HTML pages').font = BB
l4.cell(row=tr, column=5, value=f'=SUM(E2:E{tr - 1})').font = BB
l4.cell(row=tr + 2, column=1, value='x-default hreflang points to the English (Global) version. URL slugs are identical in every language.').font = Font(name=FONT, size=9, italic=True)

# Sheet 5: Redirects
r5 = wb.create_sheet('301 Redirects')
header(r5, ['#', 'Old URL', 'New URL', 'Reason'], [6, 56, 56, 46])
reason = {'/installation-and-commissioning/': 'Old duplicate page', '/wheel-shop-automation/': "Merged into Wheel Shops at HYT's request"}
for i, (o, n) in enumerate(R.REDIRECTS + R.EXTRA_REDIRECTS, 1):
    for j, v in enumerate([i, BASE + o, BASE + n, reason.get(o, 'Workbook Redirects tab (old URL)')], 1):
        c = r5.cell(row=i + 1, column=j, value=v); c.font = B; c.border = BOX
rr = len(R.REDIRECTS + R.EXTRA_REDIRECTS) + 3
r5.cell(row=rr, column=1, value='Each redirect also applies inside every language folder (e.g. /de-de/epc-project/ → /de-de/turnkey-project/). Must be server-side 301s (.htaccess, Netlify _redirects or nginx – files provided).').font = Font(name=FONT, size=9, italic=True)

# Sheet 0: Summary (first)
s = wb.create_sheet('Summary', 0)
s.column_dimensions['A'].width = 44; s.column_dimensions['B'].width = 16; s.column_dimensions['C'].width = 60
s['A1'] = 'HYT Engineering Company Pvt. Ltd. – Website Page List'; s['A1'].font = Font(name=FONT, size=14, bold=True, color='184A72')
s['A2'] = 'Prepared for vendor scoping. Generated from the website build (src/registry.py + English content).'; s['A2'].font = Font(name=FONT, size=9, italic=True)
s['A4'], s['B4'], s['C4'] = 'Item', 'Count', 'Note'
for c in ('A4', 'B4', 'C4'): s[c].font, s[c].fill, s[c].border = H, HF, BOX
PL = f"'Page List'!"
items = [
    ('Unique pages (per language)', f'=COUNTA({PL}A2:A{last})', 'One row per page in "Page List"'),
    ('Language versions', "=COUNTA(Languages!A2:A7)", 'See "Languages"'),
    ('Total HTML pages to build', '=B5*B6', 'Unique pages × languages'),
    ('Level 1 – Main pages', f'=COUNTIF({PL}B2:B{last},"Level 1*")', ''),
    ('Level 2 – Sub pages', f'=COUNTIF({PL}B2:B{last},"Level 2*")', ''),
    ('Level 3 – Sub-sub pages', f'=COUNTIF({PL}B2:B{last},"Level 3*")', ''),
    ('Product pages (shared product template)', f'=COUNTIF({PL}G2:G{last},"Product page")', 'Specs table, video, gated catalogue, sub-navigation'),
    ('Content pages', f'=COUNTIF({PL}G2:G{last},"Content page")', ''),
    ('Legal pages', f'=COUNTIF({PL}G2:G{last},"Legal page")', ''),
    ('Utility pages (thank-you, sitemap, 404 etc.)', f'=COUNTIF({PL}G2:G{last},"Utility page")', ''),
    ('Pages with forms', f'=COUNTIF({PL}J2:J{last},"*form*")', 'Enquiry, career, supplier, catalogue'),
    ('Pages with content pending HYT approval', f'=COUNTIF({PL}N2:N{last},"Yes")', 'Yellow review note on the preview site'),
    ('301 redirects (root + every language folder)', "=COUNTA('301 Redirects'!B2:B30)*B6", 'See "301 Redirects"'),
]
for i, (a, b, c) in enumerate(items, 5):
    s.cell(row=i, column=1, value=a).font = B
    s.cell(row=i, column=2, value=b).font = BB
    s.cell(row=i, column=3, value=c).font = B
    for j in (1, 2, 3): s.cell(row=i, column=j).border = BOX
n = 5 + len(items) + 1
s.cell(row=n, column=1, value='Legend (Page List)').font = BB
for k, (lbl, col) in enumerate([('Level 0 – Home', 'F58220'), ('Level 1 – Main page', 'DCE8F3'), ('Level 2 – Sub page', 'FFFFFF'), ('Level 3 – Sub-sub page', 'F7F9FB')], 1):
    c = s.cell(row=n + k, column=1, value=lbl); c.font = B; c.fill = PatternFill('solid', fgColor=col); c.border = BOX
s.cell(row=n + 6, column=1, value='Column "Vendor Remarks" in Page List is for the vendor to fill in (questions, estimates, comments).').font = Font(name=FONT, size=9, italic=True)
s.cell(row=n + 7, column=1, value='Global features: mega menu, site search, language switcher, sets no cookies, responsive, WCAG 2.1 AA, Lenis smooth scroll, animations (reduced-motion safe).').font = Font(name=FONT, size=9, italic=True)

for w in wb.worksheets:
    for row in w.iter_rows():
        for c in row:
            if c.font is None or c.font.name != FONT:
                c.font = Font(name=FONT, size=c.font.size if c.font else 10, bold=c.font.bold if c.font else False,
                              italic=c.font.italic if c.font else False, color=c.font.color if c.font else None)
out = sys.argv[1] if len(sys.argv) > 1 else str(ROOT / "docs" / "HYT-Website-Page-List.xlsx")
wb.save(out)
print('rows', len(rows), 'saved', out)
